from io import BytesIO

from docx import Document
from openpyxl import load_workbook
import fitz
import pytest
from pydantic import ValidationError

from app.api.routes_reports import ReportRequest
from app.services.reporting import export_docx, export_pdf, export_xlsx, validate_report


def _report():
    document_id = "9b6fff46-06a9-48c6-938b-83ae0cb419c7"
    run_id = "3b143e45-3beb-41f0-8702-9e6857434580"
    return {
        "title": "Document Intelligence Report",
        "report_type": "DOCUMENT_INTELLIGENCE",
        "generated_at": "2026-09-26T00:00:00+00:00",
        "parameters": {"scope": "existing indexed corpus", "document_ids": [document_id]},
        "provenance": {"source_document_ids": [document_id], "intelligence_run_id": run_id},
        "sections": [{"title": "Document Overview", "text": f"Deterministic Step 12 intelligence derived from 18189 valid indexed chunks. Intelligence run {run_id}.", "citations": []},
                      {"title": "Topics", "status": "OK", "text": "Observed from persisted intelligence.",
                       "items": [{"topic": "coal production", "frequency": 5,
                                  "document_frequencies": {document_id: 1}, "citations": ["E1"]}], "citations": ["E1"]},
                      {"title": "Keywords", "status": "OK", "items": [], "citations": []}],
        "evidence": [{"citation_id": "E1", "document_id": "document", "document_name": "real-source.pdf", "document_version_id": "version",
                      "page_number": 2, "source_unit_id": "unit", "chunk_id": "chunk", "evidence_type": "SOURCE_DOCUMENT",
                      "text": "Evidence <with> ampersand & context."}],
        "limitations": ["No production metrics supplied."],
        "validation": {"status": "VALID"},
    }


def test_report_validation_requires_resolvable_unique_citations():
    report = _report()
    assert validate_report(report)["status"] == "VALID"
    report["sections"][0]["citations"] = ["E9"]
    assert validate_report(report)["status"] == "VALIDATION_FAILED"
    report["sections"][0]["citations"] = ["E1"]
    report["evidence"].append(dict(report["evidence"][0]))
    assert "unique" in validate_report(report)["errors"][0]


def test_insufficient_verified_data_is_reported_without_invented_values():
    report = _report()
    report["report_type"] = "PRODUCTION"
    report["evidence"] = []
    report["sections"] = [{"title": "Production Overview", "status": "INSUFFICIENT_VERIFIED_DATA",
                            "text": "No verified production records match."}]
    assert validate_report(report)["status"] == "INSUFFICIENT_DATA"
    assert "value" not in report["sections"][0]


def test_pdf_export_opens_and_contains_report_content():
    content = export_pdf(_report())
    assert content.startswith(b"%PDF") and len(content) > 500
    pdf = fitz.open(stream=content, filetype="pdf")
    text = " ".join(page.get_text() for page in pdf)
    assert "Document Intelligence Report" in text
    assert "coal production" in text
    assert "real-source.pdf" in text
    headings = ["1. EXECUTIVE SUMMARY", "2. DOCUMENT ANALYSIS", "3. KEY FINDINGS",
                "4. DATA QUALITY", "APPENDIX - EVIDENCE & PROVENANCE"]
    assert [text.index(heading) for heading in headings] == sorted(text.index(heading) for heading in headings)
    assert "4. EVIDENCE" not in text
    assert "Evidence <with> ampersand & context." in text
    assert "The source material references" not in text
    assert len(pdf) <= 2
    page_one = pdf[0].get_text()
    page_two = pdf[1].get_text()
    assert "4. DATA QUALITY" in page_one
    assert "APPENDIX - EVIDENCE & PROVENANCE" not in page_one
    assert "TECHNICAL PROVENANCE & AUDIT METADATA" not in page_one
    page_two_content = [line for line in page_two.splitlines() if line.strip() not in {"DATA MINE", "Page 2"}]
    assert page_two_content[0] == "APPENDIX - EVIDENCE & PROVENANCE"
    assert "real-source.pdf" in page_two and "E1" in page_two
    for internal_value in ("document_frequencies", "chunk_id", "source_unit_id",
                           "9b6fff46-06a9-48c6-938b-83ae0cb419c7",
                           "3b143e45-3beb-41f0-8702-9e6857434580"):
        assert internal_value.lower() not in text.lower()


def test_docx_export_is_compact_and_keeps_provenance_at_the_end():
    content = export_docx(_report())
    assert len(content) > 1000
    document = Document(BytesIO(content))
    text = " ".join(node.text or "" for node in document.element.body.iter() if node.tag.endswith("}t"))
    assert "Document Intelligence Report" in text
    headings = ["1. EXECUTIVE SUMMARY", "2. DOCUMENT ANALYSIS", "3. KEY FINDINGS",
                "4. DATA QUALITY", "5. TECHNICAL PROVENANCE & AUDIT METADATA"]
    assert [text.index(heading) for heading in headings] == sorted(text.index(heading) for heading in headings)
    assert "real-source.pdf" in text
    assert "4. EVIDENCE" not in text
    assert "Evidence <with> ampersand & context." not in text
    assert "The source material references" not in text
    finding_heading = next(i for i, paragraph in enumerate(document.paragraphs) if paragraph.text == "3. KEY FINDINGS")
    quality_heading = next(i for i, paragraph in enumerate(document.paragraphs) if paragraph.text == "4. DATA QUALITY")
    assert len([paragraph for paragraph in document.paragraphs[finding_heading + 1:quality_heading] if paragraph.text.strip()]) <= 2
    technical_at = text.index("5. TECHNICAL PROVENANCE & AUDIT METADATA")
    for internal_value in ("Topics", "Keywords", "Frequency", "document_frequencies", "document_ids",
                           "intelligence run", "chunk_id", "9b6fff46-06a9-48c6-938b-83ae0cb419c7",
                           "3b143e45-3beb-41f0-8702-9e6857434580"):
        assert internal_value.lower() not in text[:technical_at].lower()
    assert "3b143e45-3beb-41f0-8702-9e6857434580" in text[technical_at:]


def test_docx_technical_provenance_is_final_and_compact():
    report = _report()
    marker = "3b143e45-3beb-41f0-8702-9e6857434580"
    report["zz_large_audit_value"] = ("x" * 30100) + marker
    document = Document(BytesIO(export_docx(report)))
    document_text = " ".join(node.text or "" for node in document.element.body.iter() if node.tag.endswith("}t"))
    technical_at = document_text.index("5. TECHNICAL PROVENANCE & AUDIT METADATA")
    assert marker not in document_text[:technical_at]
    assert marker in document_text[technical_at:]
    assert "zz_large_audit_value" not in document_text
    assert "Evidence <with> ampersand & context." not in document_text


def test_xlsx_export_opens_with_summary_sources_and_data():
    content = export_xlsx(_report())
    workbook = load_workbook(BytesIO(content), data_only=True)
    assert workbook.sheetnames == ["EXECUTIVE SUMMARY", "DOCUMENT ANALYSIS", "KEY FINDINGS", "DATA QUALITY", "TECHNICAL PROVENANCE"]
    findings_text = " ".join(str(cell.value) for row in workbook["KEY FINDINGS"].iter_rows() for cell in row if cell.value)
    assert "coal production" in findings_text
    assert "The source material references" not in findings_text
    assert workbook["KEY FINDINGS"].max_row <= 3
    assert workbook["DOCUMENT ANALYSIS"]["B2"].value == "real-source.pdf"
    main_text = " ".join(str(cell.value) for sheet in workbook.worksheets[:-1] for row in sheet.iter_rows() for cell in row if cell.value)
    for internal_value in ("document_frequencies", "document_ids", "intelligence run", "chunk_id",
                           "9b6fff46-06a9-48c6-938b-83ae0cb419c7", "3b143e45-3beb-41f0-8702-9e6857434580"):
        assert internal_value.lower() not in main_text.lower()
    technical_text = " ".join(str(cell.value) for row in workbook["TECHNICAL PROVENANCE"].iter_rows() for cell in row if cell.value)
    assert "document_frequencies" in technical_text
    assert "document_ids" in technical_text
    assert "3b143e45-3beb-41f0-8702-9e6857434580" in technical_text
    assert "Evidence payload" in technical_text
    assert "Evidence <with> ampersand & context." not in technical_text
    quality_text = " ".join(str(cell.value or "") for row in workbook["DATA QUALITY"].iter_rows() for cell in row)
    for item in ("Validation status", "Verified/trusted data availability", "Missing document data",
                 "Unresolved extraction", "Limitations", "Contradictions"):
        assert item in quality_text


def test_xlsx_export_neutralizes_formula_like_source_values():
    report = _report()
    report["evidence"][0]["document_name"] = "=HYPERLINK(\"https://bad.invalid\")"
    workbook = load_workbook(BytesIO(export_xlsx(report)), data_only=False)
    assert workbook["DOCUMENT ANALYSIS"]["B2"].data_type == "s"
    assert workbook["DOCUMENT ANALYSIS"]["B2"].value.startswith("'")


def test_report_configuration_rejects_client_supplied_business_values():
    with pytest.raises(ValidationError):
        ReportRequest(report_type="PRODUCTION", title="No client-authored numbers", production_value=42)


def test_multi_document_configuration_allows_one_real_source_for_insufficient_set_result():
    request = ReportRequest(report_type="MULTI_DOCUMENT_COMPARISON", title="Check distinct source set",
                            document_ids=["9b6fff46-06a9-48c6-938b-83ae0cb419c7"])
    assert len(request.document_ids) == 1
