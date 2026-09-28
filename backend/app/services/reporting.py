"""Deterministic report composition and export helpers."""
from __future__ import annotations

import json
import numbers
import re
from datetime import datetime, timezone
from hashlib import sha256
from html import escape
from io import BytesIO
from pathlib import PurePath


_UUID = re.compile(r"\b[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\b", re.I)
_INTERNAL_TEXT = re.compile(
    r"\b(document_ids?|version_ids?|chunk_ids?|source_unit_ids?|document_frequencies|"
    r"intelligence run|generation parameters)\b", re.I
)
_CITATION = re.compile(r"E[0-9]+\Z")
_DOCUMENT_SECTIONS = {"document overview", "executive summary", "topics", "keywords"}
_TRUSTED_EVIDENCE_TYPES = {"TRUSTED_PRODUCTION_RECORD", "TRUSTED_SPATIAL_ENTITY", "VERIFIED_CANONICAL_RECORD"}


def validate_report(report: dict) -> dict:
    problems = []
    if not report.get("title") or not report.get("sections"):
        problems.append("Report title and at least one section are required")
    required_sections = {
        "PRODUCTION": {"Production Overview"},
        "DOCUMENT_INTELLIGENCE": {"Document Overview", "Topics", "Keywords"},
        "MULTI_DOCUMENT_COMPARISON": {"Multi-document Comparison"},
        "EXECUTIVE_SUMMARY": {"Executive Summary"},
    }
    titles = {section.get("title") for section in report.get("sections", [])}
    missing = required_sections.get(report.get("report_type"), set()) - titles
    if missing:
        problems.append("Required report section missing: " + ", ".join(sorted(missing)))
    citation_ids = [item.get("citation_id") for item in report.get("evidence", [])]
    if len(citation_ids) != len(set(citation_ids)):
        problems.append("Evidence citation IDs must be unique")
    valid_refs = set(citation_ids)
    for section in report.get("sections", []):
        for ref in section.get("citations", []):
            if ref not in valid_refs:
                problems.append(f"Unresolved citation: {ref}")
        if section.get("status") == "INSUFFICIENT_VERIFIED_DATA" and section.get("items"):
            problems.append("Insufficient verified data section cannot contain reported values")
    for item in report.get("evidence", []):
        kind = item.get("evidence_type", "SOURCE_DOCUMENT")
        if kind == "SOURCE_DOCUMENT" and not all(item.get(field) for field in ("document_id", "document_version_id", "chunk_id")):
            problems.append("Source document evidence is missing document/version/chunk provenance")
        if kind == "TRUSTED_PRODUCTION_RECORD" and not all(item.get(field) for field in ("record_id", "verification_id")):
            problems.append("Trusted production evidence is missing record/verification provenance")
        if kind == "TRUSTED_SPATIAL_ENTITY" and not all(item.get(field) for field in ("entity_id", "provenance")):
            problems.append("Trusted spatial evidence is missing entity provenance")
    insufficient = any(s.get("status") in {"INSUFFICIENT_VERIFIED_DATA", "INSUFFICIENT_DOCUMENT_SET", "INSUFFICIENT_EVIDENCE", "INSUFFICIENT_SPATIAL_DATA"}
                       for s in report.get("sections", []))
    return {"status": "VALIDATION_FAILED" if problems else "INSUFFICIENT_DATA" if insufficient else "VALID",
            "errors": problems, "evidence_count": len(citation_ids), "section_count": len(report.get("sections", []))}


def _display_value(value) -> str:
    if value is None:
        return ""
    if isinstance(value, dict):
        return "; ".join(f"{key}: {_display_value(item)}" for key, item in value.items())
    if isinstance(value, (list, tuple)):
        return ", ".join(_display_value(item) for item in value)
    return str(value)


def _public_text(value) -> str:
    """Remove internal identifiers from any text rendered to a reader."""
    text = _UUID.sub("[identifier omitted]", _display_value(value)).strip()
    return _INTERNAL_TEXT.sub("[internal field omitted]", text)


def _safe_narrative(value) -> str:
    text = _display_value(value).strip()
    return "" if _UUID.search(text) or _INTERNAL_TEXT.search(text) else text


def _citation(value) -> str:
    rendered = _display_value(value).strip()
    return rendered if _CITATION.fullmatch(rendered) else ""


def _source_names(report: dict) -> list[str]:
    names = []
    for item in report.get("evidence", []):
        name = _safe_narrative(item.get("document_name"))
        if name and name not in names:
            names.append(name)
    return names


def _metadata_records(report: dict) -> list[dict]:
    records = []
    for key in ("document_metadata", "source_documents", "documents", "metadata"):
        value = report.get(key)
        if isinstance(value, list):
            records.extend(row for row in value if isinstance(row, dict))
        elif isinstance(value, dict):
            if any(field in value for field in ("document_name", "original_filename", "page_count", "ocr_status")):
                records.append(value)
            else:
                records.extend(row for row in value.values() if isinstance(row, dict))
    return records


def _document_metadata_value(report: dict, name: str, fields: tuple[str, ...]) -> str:
    for field in fields:
        if report.get(field) not in (None, ""):
            return _public_text(report[field])
    for row in _metadata_records(report):
        row_name = next((row.get(key) for key in ("document_name", "original_filename", "name") if row.get(key)), None)
        if row_name and str(row_name) != name:
            continue
        for field in fields:
            if row.get(field) not in (None, ""):
                return _public_text(row[field])
    for section in report.get("sections", []):
        for row in (section.get("items") or []):
            if not isinstance(row, dict):
                continue
            row_name = row.get("document_name") or row.get("original_filename")
            if row_name and str(row_name) != name:
                continue
            for field in fields:
                if row.get(field) not in (None, ""):
                    return _public_text(row[field])
    for evidence in report.get("evidence", []):
        if evidence.get("document_name") == name:
            for field in fields:
                if evidence.get(field) not in (None, ""):
                    return _public_text(evidence[field])
    return ""


def _readable_row(row: dict) -> dict:
    internal = {
        "citation", "citations", "document_frequencies", "provenance", "metadata",
        "document_id", "document_version_id", "version_id", "chunk_id", "source_unit_id",
        "record_id", "verification_id", "entity_id", "id",
    }
    result = {}
    for key, value in row.items():
        key_text = str(key).lower()
        if key_text in internal or key_text.endswith("_id") or key_text.startswith("_"):
            continue
        if "frequenc" in key_text:
            continue
        if value in (None, "", [], {}):
            continue
        safe_value = _public_text(value)
        if safe_value:
            result[str(key).replace("_", " ").strip().title()] = safe_value
    return result


def _group_topic_names(topic_names: list[str]) -> list[str]:
    """Create concise topic groups using only labels present in the report payload."""
    remaining = list(topic_names)
    groups = []

    def add_group(prefix: str, patterns: tuple[str, ...]) -> None:
        matched = [name for name in remaining if any(pattern in name.lower() for pattern in patterns)]
        if matched:
            shown = matched[:5]
            suffix = f" and {len(matched) - len(shown)} related labels" if len(matched) > len(shown) else ""
            groups.append(prefix + ", ".join(shown) + suffix)
            for name in matched:
                remaining.remove(name)

    add_group("coal types and products (", ("raw coal", "coking coal", "non-coking coal", "coal coke"))
    if groups and groups[-1].startswith("coal types and products ("):
        groups[-1] += ")"
    add_group("geological resources (", ("geological resource", "coal resource", "resources", "geology"))
    if groups and groups[-1].startswith("geological resources ("):
        groups[-1] += ")"
    add_group("production and dispatch (", ("production", "dispatch"))
    if groups and groups[-1].startswith("production and dispatch ("):
        groups[-1] += ")"
    add_group("stocks (", ("closing stock", "pit-head stock", "stock", "inventory"))
    if groups and groups[-1].startswith("stocks ("):
        groups[-1] += ")"
    add_group("pricing, royalty and taxes (", ("price", "pricing", "royalty", "tax"))
    if groups and groups[-1].startswith("pricing, royalty and taxes ("):
        groups[-1] += ")"
    add_group("captive/commercial coal blocks (", ("coal block", "captive", "commercial block"))
    if groups and groups[-1].startswith("captive/commercial coal blocks ("):
        groups[-1] += ")"
    add_group("international statistics (", ("international", "world coal", "global coal"))
    if groups and groups[-1].startswith("international statistics ("):
        groups[-1] += ")"
    if remaining:
        groups.append("other analyzed areas (" + ", ".join(remaining[:5])
                      + (f" and {len(remaining) - 5} related labels" if len(remaining) > 5 else "") + ")")
    return groups


def _compact_source_references(report: dict) -> str:
    refs = []
    for section in report.get("sections", []):
        refs.extend(_citation(ref) for ref in section.get("citations", []))
        for row in section.get("items", []) or []:
            if isinstance(row, dict):
                refs.extend(_citation(ref) for ref in row.get("citations", []))
    refs = list(dict.fromkeys(ref for ref in refs if ref))
    if not refs:
        return ""
    evidence_by_ref = {
        ref: item for item in report.get("evidence", [])
        if (ref := _citation(item.get("citation_id")))
    }
    pages_by_document = {}
    for ref in refs:
        item = evidence_by_ref.get(ref)
        if not item:
            continue
        document = _safe_narrative(item.get("document_name")) or "Source document"
        page = _public_text(item.get("page_number"))
        pages = pages_by_document.setdefault(document, [])
        if page and page not in pages:
            pages.append(page)
    page_text = []
    for document, pages in pages_by_document.items():
        display_pages = pages[:6]
        suffix = f" (+{len(pages) - len(display_pages)} more)" if len(pages) > len(display_pages) else ""
        page_text.append(f"{document}, pp. {', '.join(display_pages)}{suffix}")
    display_refs = refs[:6]
    ref_text = ", ".join(display_refs)
    if len(refs) > len(display_refs):
        ref_text += f" and {len(refs) - len(display_refs)} more citations"
    location_text = "; ".join(page_text)
    return f"Supporting references: {location_text}; citations {ref_text}." if location_text else f"Supporting citations: {ref_text}."


def _quantitative_findings(report: dict) -> list[str]:
    numeric_fields = ("value", "quantity", "amount", "production", "tonnage", "grade", "verified_count", "total", "average", "median", "rate")
    findings = []
    for section in report.get("sections", []):
        title = str(section.get("title", "")).strip()
        if title.lower() in _DOCUMENT_SECTIONS or str(section.get("status", "")).startswith("INSUFFICIENT"):
            continue
        section_refs = section.get("citations", [])
        for row in section.get("items", []) or []:
            if not isinstance(row, dict):
                continue
            refs = [ref for ref in (_citation(item) for item in row.get("citations", section_refs)) if ref]
            for key, value in row.items():
                key_text = str(key).lower()
                if "frequenc" in key_text or not any(field in key_text for field in numeric_fields):
                    continue
                if not isinstance(value, numbers.Number) or isinstance(value, bool):
                    continue
                descriptors = [
                    f"{str(field).replace('_', ' ')} {_public_text(row[field])}"
                    for field in ("commodity", "mine", "project", "state", "period", "unit")
                    if row.get(field) not in (None, "") and _public_text(row[field])
                ]
                fact = f"{title}: " + (", ".join(descriptors) + ", " if descriptors else "")
                fact += f"{str(key).replace('_', ' ')} {_public_text(value)}"
                if refs:
                    fact += f" (citations {', '.join(refs[:3])})"
                findings.append(fact)
                if len(findings) >= 4:
                    return findings
    return findings


def _technical_metadata(report: dict) -> list[tuple[str, str]]:
    """Keep bulky provenance outside the narrative sections without repeating excerpts."""
    provenance = report.get("provenance") or {}
    sections = report.get("sections", [])
    evidence = report.get("evidence", [])
    overview = next((section for section in sections if str(section.get("title", "")).lower() == "document overview"), {})
    overview_provenance = overview.get("provenance") or {}
    topic_distributions = []
    for section in sections:
        if str(section.get("title", "")).lower() == "topics":
            topic_distributions = [
                {"topic": row.get("topic"), "document_frequencies": row.get("document_frequencies")}
                for row in (section.get("items") or []) if row.get("document_frequencies")
            ]
    source_ids = list(dict.fromkeys(
        [str(value) for value in provenance.get("source_document_ids", [])]
        + [str(item.get("document_id")) for item in evidence if item.get("document_id")]
    ))
    version_ids = list(dict.fromkeys(
        [str(value) for value in provenance.get("source_version_ids", [])]
        + [str(item.get("document_version_id")) for item in evidence if item.get("document_version_id")]
    ))
    chunks = list(dict.fromkeys(str(item.get("chunk_id")) for item in evidence if item.get("chunk_id")))
    source_units = list(dict.fromkeys(str(item.get("source_unit_id")) for item in evidence if item.get("source_unit_id")))
    page_ids = list(dict.fromkeys(str(item.get("page_id")) for item in evidence if item.get("page_id")))
    citation_ids = list(dict.fromkeys(ref for ref in (_citation(item.get("citation_id")) for item in evidence) if ref))
    parameters = report.get("parameters") or {}
    validation = report.get("validation") or {}
    return [
        ("Report type", _display_value(report.get("report_type"))),
        ("Generated at", _display_value(report.get("generated_at"))),
        ("Validation status", _display_value(validation.get("status"))),
        ("Generation parameters", json.dumps(parameters, ensure_ascii=False, sort_keys=True, default=str)),
        ("Source document IDs", ", ".join(source_ids) or _display_value(provenance.get("source_document_ids"))),
        ("Source version IDs", ", ".join(version_ids) or _display_value(provenance.get("source_version_ids"))),
        ("Intelligence run ID", _display_value(provenance.get("intelligence_run_id") or overview_provenance.get("intelligence_run_id"))),
        ("Source chunk count", _display_value(overview_provenance.get("source_chunk_count"))),
        ("Page IDs", ", ".join(page_ids)),
        ("Source-unit IDs", ", ".join(source_units)),
        ("Chunk IDs", ", ".join(chunks)),
        ("Citation IDs", ", ".join(citation_ids)),
        ("Document frequencies", json.dumps(topic_distributions, ensure_ascii=False, sort_keys=True, default=str)),
        ("Evidence payload", f"{len(evidence)} evidence records, including excerpts and provenance, remain attached to the report record and API response."),
    ]


def _presentation(report: dict) -> dict:
    sections = report.get("sections", [])
    by_title = {str(section.get("title", "")).strip().lower(): section for section in sections}
    evidence = report.get("evidence", [])
    names = _source_names(report)
    topic_rows = (by_title.get("topics") or {}).get("items", [])
    topic_names = list(dict.fromkeys(
        term for term in (_safe_narrative(row.get("topic")) for row in topic_rows) if term
    ))
    groups = _group_topic_names(topic_names)
    source_text = ", ".join(names) if names else "the supplied source material"
    summary = (
        f"This report summarizes the available source-backed analysis for {source_text}. "
        "Supporting evidence, citations, and provenance remain attached to the report record."
    )

    areas = "; ".join(groups) if groups else "No major content-area labels were included in the report payload."
    document_rows = []
    for name in names or ["Not available in report payload"]:
        file_type = _document_metadata_value(report, name, ("document_type", "mime_type", "content_type", "file_type"))
        if not file_type and name != "Not available in report payload":
            suffix = PurePath(name).suffix.lower().lstrip(".")
            file_type = suffix.upper() if suffix else "Not available in report payload"
        document_rows.append({
            "Source document": name,
            "Document type": file_type or "Not available in report payload",
            "Page count": _document_metadata_value(report, name, ("page_count", "total_pages")) or "Not available in report payload",
            "Processing status": _document_metadata_value(report, name, ("processing_status", "document_status")) or "Not available in report payload",
            "OCR status": _document_metadata_value(report, name, ("ocr_status", "ocr_state")) or "Not available in report payload",
            "Major content areas": areas,
        })

    findings = []
    if groups:
        findings.append("The indexed content covers " + "; ".join(groups) + ".")
    quantitative = _quantitative_findings(report)
    if quantitative:
        findings.append("Source-backed quantitative results in the report payload include " + "; ".join(quantitative[:4]) + ".")
    references = _compact_source_references(report)
    if references:
        if findings:
            findings[-1] += " " + references
        else:
            findings.append("Source-backed findings are available in the report payload. " + references)
    if not findings:
        findings.append("No source-backed finding summary was included in the report payload.")

    quality = []
    validation = report.get("validation") or {}
    quality.append({"Item": "Validation status", "Result": _public_text(validation.get("status", "Not available in report payload"))})
    trusted_count = sum(1 for item in evidence if str(item.get("evidence_type", "")).upper() in _TRUSTED_EVIDENCE_TYPES)
    trusted_result = (
        f"{trusted_count} trusted-record evidence reference(s) are included."
        if trusted_count else "No verified/trusted record evidence is included; the available evidence is source-document text."
    )
    quality.append({"Item": "Verified/trusted data availability", "Result": trusted_result})
    missing_fields = []
    for _, key in (("page count", "Page count"), ("processing status", "Processing status"), ("OCR status", "OCR status")):
        if not document_rows or any(row.get(key) == "Not available in report payload" for row in document_rows):
            missing_fields.append(key.lower())
    quality.append({"Item": "Missing document data", "Result": ("Not included in report payload: " + ", ".join(missing_fields)) if missing_fields else "No missing document metadata fields were identified in the report payload."})
    unresolved = report.get("unresolved_extraction")
    if unresolved is None:
        unresolved = report.get("unresolved_extraction_count")
    if unresolved is None:
        unresolved_section = next((section for section in sections if "unresolved extraction" in str(section.get("title", "")).lower()), None)
        unresolved = _safe_narrative(unresolved_section.get("text")) if unresolved_section else "Not included in report payload."
    quality.append({"Item": "Unresolved extraction", "Result": _public_text(unresolved)})

    limitations = []
    for value in report.get("limitations", []):
        safe = _safe_narrative(value)
        if safe and safe not in limitations:
            limitations.append(safe)
    for section in sections:
        status = str(section.get("status", ""))
        if status in {"WARNING", "INSUFFICIENT_VERIFIED_DATA", "INSUFFICIENT_DOCUMENT_SET", "INSUFFICIENT_EVIDENCE", "INSUFFICIENT_SPATIAL_DATA"}:
            title = _safe_narrative(section.get("title")) or "Report section"
            text = _safe_narrative(section.get("text"))
            item = f"{title}: {status}" + (f" — {text}" if text else "")
            if item not in limitations:
                limitations.append(item)
    quality.append({"Item": "Limitations", "Result": "; ".join(limitations) if limitations else "No limitations were included in the report payload."})

    contradictions = report.get("contradictions")
    if contradictions is None:
        contradiction_section = next((section for section in sections if any(word in str(section.get("title", "")).lower() for word in ("contradiction", "conflict"))), None)
        contradictions = (contradiction_section.get("text") or contradiction_section.get("items") or contradiction_section.get("status")) if contradiction_section else "No contradiction assessment was included in the report payload."
    quality.append({"Item": "Contradictions", "Result": _public_text(contradictions)})

    return {"summary": summary, "documents": document_rows, "findings": findings,
            "quality": quality, "technical": _technical_metadata(report)}


def _group_topic_names(topic_names: list[str]) -> list[str]:
    """Group only topic labels that are actually present in the report payload."""
    remaining = list(topic_names)
    groups = []

    def take(label: str, patterns: tuple[str, ...], prefix: str = "") -> None:
        found = [name for name in remaining if any(pattern in name.lower() for pattern in patterns)]
        if found:
            groups.append(prefix + ", ".join(found[:5]))
            for name in found:
                remaining.remove(name)

    take("", ("raw coal", "coking coal", "non-coking coal", "coal coke"), "coal types/products: ")
    take("", ("geological resource", "coal resource", "resources", "geology"), "resources/geology: ")
    take("", ("production", "dispatch"), "production/dispatch: ")
    take("", ("closing stock", "pit-head stock", "stock", "inventory"), "stocks: ")
    take("", ("price", "pricing", "royalty", "tax"), "pricing/royalty/taxes: ")
    take("", ("coal block", "captive", "commercial block"), "captive/commercial blocks: ")
    take("", ("international", "world coal", "global coal"), "international statistics: ")
    groups.extend(remaining[:5])
    return groups


def _compact_source_references(report: dict) -> str:
    sections = report.get("sections", [])
    refs = []
    for section in sections:
        refs.extend(_citation(ref) for ref in section.get("citations", []))
        for row in section.get("items", []) or []:
            if isinstance(row, dict):
                refs.extend(_citation(ref) for ref in row.get("citations", []))
    refs = list(dict.fromkeys(ref for ref in refs if ref))
    if not refs:
        return ""
    lookup = {ref: item for item in report.get("evidence", []) if (ref := _citation(item.get("citation_id")))}
    page_groups = {}
    for ref in refs:
        item = lookup.get(ref)
        if not item:
            continue
        name = _safe_narrative(item.get("document_name")) or "Source document"
        page = _public_text(item.get("page_number"))
        if page and page not in page_groups.setdefault(name, []):
            page_groups[name].append(page)
    source_labels = []
    for name, pages in page_groups.items():
        shown = pages[:8]
        suffix = f" +{len(pages) - len(shown)} more" if len(pages) > len(shown) else ""
        source_labels.append(f"{name}, pp. {', '.join(shown)}{suffix}")
    shown_refs = refs[:8]
    ref_text = ", ".join(shown_refs)
    if len(refs) > len(shown_refs):
        ref_text += f" +{len(refs) - len(shown_refs)} more"
    return "Supporting sources: " + "; ".join(source_labels) + f" (citations {ref_text})."


def _quantitative_findings(report: dict) -> list[str]:
    accepted = ("value", "quantity", "amount", "production", "tonnage", "grade", "verified_count", "total", "average", "median", "rate")
    findings = []
    for section in report.get("sections", []):
        title = str(section.get("title", "")).strip()
        if title.lower() in _DOCUMENT_SECTIONS or str(section.get("status", "")).startswith("INSUFFICIENT"):
            continue
        section_refs = [_citation(ref) for ref in section.get("citations", [])]
        for row in section.get("items", []) or []:
            if not isinstance(row, dict):
                continue
            refs = [_citation(ref) for ref in row.get("citations", section_refs)]
            refs = [ref for ref in refs if ref]
            for key, value in row.items():
                key_lower = str(key).lower()
                if not any(token in key_lower for token in accepted) or "frequenc" in key_lower:
                    continue
                if not isinstance(value, numbers.Number) or isinstance(value, bool):
                    continue
                dimensions = [
                    f"{str(dim).replace('_', ' ')} {_public_text(row[dim])}"
                    for dim in ("commodity", "mine", "project", "state", "period", "unit")
                    if row.get(dim) not in (None, "") and _public_text(row[dim])
                ]
                label = ", ".join(dimensions)
                fact = f"{title}: " + (label + ", " if label else "") + f"{str(key).replace('_', ' ')} {_public_text(value)}"
                if refs:
                    fact += f" [{', '.join(refs[:3])}]"
                findings.append(fact)
                if len(findings) >= 4:
                    return findings
    return findings


def _technical_metadata(report: dict) -> list[tuple[str, str]]:
    provenance = report.get("provenance") or {}
    evidence = report.get("evidence", [])
    sections = report.get("sections", [])
    overview = next((section for section in sections if str(section.get("title", "")).lower() == "document overview"), {})
    overview_provenance = overview.get("provenance") or {}
    topic_distributions = []
    for section in sections:
        if str(section.get("title", "")).lower() != "topics":
            continue
        topic_distributions = [
            {"topic": row.get("topic"), "document_frequencies": row.get("document_frequencies")}
            for row in (section.get("items") or []) if row.get("document_frequencies")
        ]
    document_ids = list(dict.fromkeys(
        [str(item) for item in provenance.get("source_document_ids", [])]
        + [str(item.get("document_id")) for item in evidence if item.get("document_id")]
    ))
    version_ids = list(dict.fromkeys(
        [str(item) for item in provenance.get("source_version_ids", [])]
        + [str(item.get("document_version_id")) for item in evidence if item.get("document_version_id")]
    ))
    chunks = list(dict.fromkeys(str(item.get("chunk_id")) for item in evidence if item.get("chunk_id")))
    source_units = list(dict.fromkeys(str(item.get("source_unit_id")) for item in evidence if item.get("source_unit_id")))
    page_ids = list(dict.fromkeys(str(item.get("page_id")) for item in evidence if item.get("page_id")))
    citation_ids = list(dict.fromkeys(ref for ref in (_citation(item.get("citation_id")) for item in evidence) if ref))
    parameters = report.get("parameters") or {}
    validation = report.get("validation") or {}
    return [
        ("Report type", _display_value(report.get("report_type"))),
        ("Generated at", _display_value(report.get("generated_at"))),
        ("Validation status", _display_value(validation.get("status"))),
        ("Generation parameters", json.dumps(parameters, ensure_ascii=False, sort_keys=True, default=str)),
        ("Source document IDs", ", ".join(document_ids) or _display_value(provenance.get("source_document_ids"))),
        ("Source version IDs", ", ".join(version_ids) or _display_value(provenance.get("source_version_ids"))),
        ("Intelligence run ID", _display_value(provenance.get("intelligence_run_id") or overview_provenance.get("intelligence_run_id"))),
        ("Source chunk count", _display_value(overview_provenance.get("source_chunk_count"))),
        ("Page IDs", ", ".join(page_ids)),
        ("Source-unit IDs", ", ".join(source_units)),
        ("Chunk IDs", ", ".join(chunks)),
        ("Citation IDs", ", ".join(citation_ids)),
        ("Document frequencies", json.dumps(topic_distributions, ensure_ascii=False, sort_keys=True, default=str)),
        ("Evidence payload", f"{len(evidence)} evidence records, including excerpts and provenance, remain attached to the report record and API response."),
    ]


def _pdf_table(rows, colors, Table, TableStyle, Paragraph, small):
    data = [[Paragraph(escape(str(key)), small) for key in rows[0]]]
    data.extend([[Paragraph(escape(str(value)), small) for value in row.values()] for row in rows])
    table = Table(data, repeatRows=1, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#18324B")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 5),
        ("RIGHTPADDING", (0, 0), (-1, -1), 5),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]))
    return table


def export_pdf(report: dict) -> bytes:
    """Export a concise human report followed by a readable evidence appendix."""
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, PageBreak, KeepTogether, Table, TableStyle

    view = _presentation(report)
    stream = BytesIO()
    doc = SimpleDocTemplate(stream, pagesize=A4, rightMargin=16 * mm, leftMargin=16 * mm,
                            topMargin=16 * mm, bottomMargin=17 * mm)
    base = getSampleStyleSheet()
    brand = ParagraphStyle("DMBrand", parent=base["Title"], fontSize=21, leading=25, alignment=TA_CENTER, spaceAfter=5)
    title = ParagraphStyle("DMReportTitle", parent=base["Heading1"], fontSize=15, leading=18, spaceAfter=10)
    section = ParagraphStyle("DMContentSection", parent=base["Heading2"], fontSize=13, leading=16, spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("DMContentBody", parent=base["BodyText"], fontSize=9, leading=12, spaceAfter=5)
    small = ParagraphStyle("DMContentSmall", parent=base["BodyText"], fontSize=7.5, leading=9)
    appendix_title = ParagraphStyle("DMAppendixTitle", parent=base["Heading1"], fontSize=16, leading=20, spaceAfter=8)
    evidence_text = ParagraphStyle("DMEvidenceText", parent=base["BodyText"], fontSize=8.5, leading=11, leftIndent=8, spaceAfter=7)

    all_evidence = report.get("evidence", [])
    evidence_by_citation = {}
    for item in all_evidence:
        ref = _citation(item.get("citation_id") or item.get("reference_id") or item.get("citation"))
        if ref and ref not in evidence_by_citation:
            evidence_by_citation[ref] = item

    selected_evidence = []
    seen_references = set()
    for section_data in report.get("sections", []):
        if str(section_data.get("title", "")).strip().lower() != "topics":
            continue
        for topic_row in section_data.get("items", []) or []:
            if not isinstance(topic_row, dict):
                continue
            topic_name = _safe_narrative(topic_row.get("topic"))
            citations = topic_row.get("citations") or []
            if isinstance(citations, str):
                citations = [citations]
            for raw_ref in citations:
                ref = _citation(raw_ref)
                if ref and ref in evidence_by_citation and ref not in seen_references:
                    selected_evidence.append((evidence_by_citation[ref], topic_name))
                    seen_references.add(ref)
                    break
            if len(selected_evidence) >= 10:
                break
        break
    if len(selected_evidence) < 10:
        for item in all_evidence:
            ref = _citation(item.get("citation_id") or item.get("reference_id") or item.get("citation"))
            if ref and ref not in seen_references:
                selected_evidence.append((item, ""))
                seen_references.add(ref)
            if len(selected_evidence) >= 10:
                break

    evidence_rows = []
    for item, topic_name in selected_evidence:
        citation_id = _citation(item.get("citation_id") or item.get("reference_id") or item.get("citation"))
        source_name = _safe_narrative(item.get("document_name") or item.get("original_filename"))
        page_number = _display_value(item.get("page_number") or item.get("source_page") or item.get("page")).strip()
        excerpt = _display_value(item.get("excerpt") or item.get("text") or item.get("evidence_text")).strip()
        if not (source_name or page_number or excerpt or citation_id):
            continue
        evidence_rows.append({
            "citation": citation_id or "Reference not supplied",
            "topic": topic_name,
            "source": source_name or "Source document name not supplied",
            "page": page_number or "Page not supplied",
            "excerpt": excerpt or "Excerpt not supplied",
        })

    source_names = list(dict.fromkeys(
        name for name in (_safe_narrative(item.get("document_name") or item.get("original_filename")) for item in all_evidence) if name
    ))
    unique_pages = list(dict.fromkeys(
        page for page in (_display_value(item.get("page_number") or item.get("source_page") or item.get("page")).strip() for item in all_evidence) if page
    ))
    validation = report.get("validation") or {}
    appendix_summary = [
        {"Item": "Report type", "Details": _display_value(report.get("report_type")) or "Not supplied"},
        {"Item": "Validation status", "Details": _display_value(validation.get("status")) or "Not supplied"},
        {"Item": "Evidence references", "Details": str(len(all_evidence))},
        {"Item": "Source documents represented", "Details": str(len(source_names))},
        {"Item": "Source pages represented", "Details": str(len(unique_pages))},
        {"Item": "References shown", "Details": f"{len(evidence_rows)} selected; full evidence remains available with the report record"},
    ]

    pdf_documents = []
    for index, document_row in enumerate(view["documents"], start=1):
        if len(view["documents"]) > 1:
            pdf_documents.append({"Attribute": f"Source {index}", "Details": document_row.get("Source document", "Not supplied")})
        else:
            pdf_documents.append({"Attribute": "Source document", "Details": document_row.get("Source document", "Not supplied")})
        for field, label in (("Document type", "Document type"), ("Page count", "Page count"),
                             ("Processing status", "Processing status"), ("OCR status", "OCR status")):
            if document_row.get(field):
                pdf_documents.append({"Attribute": label, "Details": document_row[field]})
        areas = document_row.get("Major content areas", "")
        area_groups = []
        for area in areas.split(";"):
            area = area.strip()
            if not area:
                continue
            label = area.split(":", 1)[0].strip() if ":" in area else "Additional reported topics"
            if label and label not in area_groups:
                area_groups.append(label)
        if areas:
            pdf_documents.append({"Attribute": "Major content areas", "Details": "; ".join(area_groups) or areas})

    pdf_findings = []
    area_labels = [row["Details"] for row in pdf_documents if row["Attribute"] == "Major content areas"]
    if area_labels:
        pdf_findings.append("The document covers " + area_labels[0] + ". Supporting source excerpts are provided in the appendix.")
    for finding in view["findings"]:
        if finding.startswith("Source-backed quantitative results"):
            pdf_findings.append(finding.split(" Supporting sources:", 1)[0])
    if not pdf_findings:
        pdf_findings = ["No concise source-backed finding summary was included in the report payload."]

    flow = [Paragraph("DATA MINE", brand),
            Paragraph(escape(_safe_narrative(report.get("title")) or "Mining Intelligence Report"), title),
            Paragraph("1. EXECUTIVE SUMMARY", section), Paragraph(escape(view["summary"]), body),
            Paragraph("2. DOCUMENT ANALYSIS", section),
            _pdf_table(pdf_documents, colors, Table, TableStyle, Paragraph, small),
            Paragraph("3. KEY FINDINGS", section)]
    flow.extend(Paragraph(escape(paragraph), body) for paragraph in pdf_findings[:2])
    flow.append(Paragraph("4. DATA QUALITY", section))
    flow.append(_pdf_table(view["quality"], colors, Table, TableStyle, Paragraph, small))

    # Keep the human-facing report together on page one; the appendix always
    # starts at the top of a fresh page and contains no raw identifier dumps.
    flow.extend([
        PageBreak(),
        Paragraph("APPENDIX - EVIDENCE &amp; PROVENANCE", appendix_title),
        Paragraph("Evidence references are shown with their source document and page where supplied. Citation IDs link to the complete evidence and provenance retained with this report.", body),
        _pdf_table(appendix_summary, colors, Table, TableStyle, Paragraph, small),
        Paragraph("Selected Evidence", section),
    ])
    if evidence_rows:
        for row in evidence_rows:
            excerpt = row["excerpt"]
            if len(excerpt) > 560:
                excerpt = excerpt[:557].rsplit(" ", 1)[0] + "..."
            topic_label = f"<b>Topic:</b> {escape(row['topic'])} &nbsp;|&nbsp; " if row["topic"] else ""
            block = [
                Paragraph(
                    f"{topic_label}<b>{escape(row['citation'])}</b> &nbsp;|&nbsp; {escape(row['source'])} &nbsp;|&nbsp; Page {escape(row['page'])}",
                    small,
                ),
                Paragraph(escape(excerpt), evidence_text),
            ]
            flow.append(KeepTogether(block))
    else:
        flow.append(Paragraph("No evidence references were included in the report payload.", body))
    flow.append(Paragraph("Technical Audit Summary", section))
    flow.append(Paragraph(
        "Internal document, run, chunk, source-unit, and topic-distribution identifiers are retained in the report's audit data and are omitted from this reader-facing PDF."
        , small
    ))

    def footer(canvas, document):
        canvas.saveState()
        canvas.setFont("Helvetica", 8)
        canvas.drawString(16 * mm, 9 * mm, "DATA MINE")
        canvas.drawRightString(A4[0] - 16 * mm, 9 * mm, f"Page {document.page}")
        canvas.restoreState()

    doc.build(flow, onFirstPage=footer, onLaterPages=footer)
    return stream.getvalue()


def export_docx(report: dict) -> bytes:
    """Export the concise report and technical provenance to Word."""
    from docx import Document

    view = _presentation(report)
    doc = Document()
    doc.add_heading("DATA MINE", 0)
    doc.add_heading(_safe_narrative(report.get("title")) or "Mining Intelligence Report", 1)
    doc.add_heading("1. EXECUTIVE SUMMARY", 2)
    doc.add_paragraph(view["summary"])
    doc.add_heading("2. DOCUMENT ANALYSIS", 2)
    document_table = doc.add_table(rows=1, cols=2)
    document_table.style = "Light Shading Accent 1"
    document_table.rows[0].cells[0].text = "Attribute"
    document_table.rows[0].cells[1].text = "Details"
    for row in view["documents"]:
        for key, value in row.items():
            cells = document_table.add_row().cells
            cells[0].text, cells[1].text = key, value
    doc.add_heading("3. KEY FINDINGS", 2)
    for finding in view["findings"]:
        doc.add_paragraph(finding, style="List Bullet")
    doc.add_heading("4. DATA QUALITY", 2)
    quality_table = doc.add_table(rows=1, cols=2)
    quality_table.style = "Light Shading Accent 1"
    quality_table.rows[0].cells[0].text, quality_table.rows[0].cells[1].text = "Item", "Result"
    for row in view["quality"]:
        cells = quality_table.add_row().cells
        cells[0].text, cells[1].text = row["Item"], row["Result"]
    doc.add_heading("5. TECHNICAL PROVENANCE & AUDIT METADATA", 2)
    technical_table = doc.add_table(rows=1, cols=2)
    technical_table.style = "Light Shading Accent 1"
    technical_table.rows[0].cells[0].text, technical_table.rows[0].cells[1].text = "Field", "Value"
    for key, value in view["technical"]:
        cells = technical_table.add_row().cells
        cells[0].text = key
        cells[1].text = value or "Not available"
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def export_xlsx(report: dict) -> bytes:
    """Export concise report sheets plus a technical provenance sheet."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    view = _presentation(report)

    def cell(value):
        rendered = _public_text(value)
        return "'" + rendered if rendered.startswith(("=", "+", "-", "@")) else rendered

    def technical_cell(value):
        # Technical provenance is intentionally retained on this sheet. Keep
        # Excel formula injection protection without redacting IDs/field names.
        rendered = _display_value(value)
        return "'" + rendered if rendered.startswith(("=", "+", "-", "@")) else rendered

    workbook = Workbook()
    executive = workbook.active
    executive.title = "EXECUTIVE SUMMARY"
    executive.append(["DATA MINE"])
    executive.append([cell(_safe_narrative(report.get("title")) or "Mining Intelligence Report")])
    executive.append([])
    executive.append(["EXECUTIVE SUMMARY"])
    executive.append([cell(view["summary"])])

    document_analysis = workbook.create_sheet("DOCUMENT ANALYSIS")
    document_analysis.append(["Attribute", "Details"])
    for row in view["documents"]:
        for key, value in row.items():
            document_analysis.append([cell(key), cell(value)])

    findings = workbook.create_sheet("KEY FINDINGS")
    findings.append(["Summary"])
    for paragraph in view["findings"]:
        findings.append([cell(paragraph)])

    quality = workbook.create_sheet("DATA QUALITY")
    quality.append(["Item", "Result"])
    for row in view["quality"]:
        quality.append([cell(row["Item"]), cell(row["Result"])])

    technical = workbook.create_sheet("TECHNICAL PROVENANCE")
    technical.append(["Field", "Value"])
    for key, value in view["technical"]:
        for start in range(0, max(len(value), 1), 30000):
            technical.append([technical_cell(key), technical_cell(value[start:start + 30000])])

    for sheet in workbook:
        sheet.freeze_panes = "A2"
        for header in sheet[1]:
            header.font = Font(bold=True, color="FFFFFF")
            header.fill = PatternFill("solid", fgColor="18324B")
            header.alignment = Alignment(vertical="top")
        for column in sheet.columns:
            letter = column[0].column_letter
            width = max((len(str(item.value or "")) for item in column), default=0) + 2
            sheet.column_dimensions[letter].width = min(70, max(14, width))
        for row in sheet.iter_rows():
            for item in row:
                item.alignment = Alignment(vertical="top", wrap_text=True)
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def artifact_metadata(report_id, kind: str, data: bytes) -> dict:
    extensions = {"PDF": ("pdf", "application/pdf"), "DOCX": ("docx", "application/vnd.openxmlformats-officedocument.wordprocessingml.document"), "XLSX": ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    ext, mime = extensions[kind]
    return {"artifact_type": kind, "filename": f"data-mine-report-{report_id}.{ext}", "mime_type": mime,
            "checksum_sha256": sha256(data).hexdigest(), "size_bytes": len(data), "data": data}


def generated_timestamp() -> str:
    return datetime.now(timezone.utc).isoformat()
