import asyncio
from decimal import Decimal
from io import BytesIO
from pathlib import Path
import shutil
from types import SimpleNamespace
from uuid import uuid4

import pymupdf
import pytest
from docx import Document as WordDocument
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image, ImageDraw
from sqlalchemy import delete, select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.orm import Session as SyncSession, sessionmaker
from sqlalchemy.pool import NullPool

from app.api.dependencies import get_storage_service
from app.core.auth import get_current_user
from app.core.config import DATABASE_URL, SYNC_DATABASE_URL
from app.db.session import get_db
from app.main import app
from app.models.documents import Document, DocumentVersion, DocumentPage
from app.models.intelligence import ExtractedContent, ExtractedTable, ExtractedTableCell
from app.models.processing import ProcessingJob, ProcessingStage
from app.processors.docx import DOCXProcessor
from app.processors.excel import ExcelProcessor
from app.processors.image import ImageProcessor
from app.processors.ocr import OCRBlock, OCRResult, OCRService, OCREngineUnavailable
from app.processors.pdf import PDFProcessor
from app.processors.router import ProcessorRouter, UnsupportedProcessorType
from app.storage.local import LocalStorageService
from app.workers.processing import process_document_job


def save_object(storage: LocalStorageService, key: str, data: bytes) -> None:
    storage.upload_file(BytesIO(data), key, "application/octet-stream")


def make_pdf(text: str = "DATA MINE processing fixture with sufficient text") -> bytes:
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), text)
    result = document.tobytes()
    document.close()
    return result


def make_table_pdf() -> bytes:
    document = pymupdf.open()
    page = document.new_page(width=400, height=300)
    for x in (50, 180, 300):
        page.draw_line((x, 60), (x, 180), width=1)
    for y in (60, 100, 140, 180):
        page.draw_line((50, y), (300, y), width=1)
    values = [["Name", "Value"], ["Alpha", "10"], ["Beta", "20"]]
    for row_index, row in enumerate(values):
        y = 85 + row_index * 40
        page.insert_text((60, y), row[0])
        page.insert_text((190, y), row[1])
    data = document.tobytes()
    document.close()
    return data


class FixedOCR:
    def recognize(self, _image_bytes: bytes) -> OCRResult:
        return OCRResult("synthetic OCR line", Decimal("0.91"), [OCRBlock("synthetic OCR line", {"x0": 1, "y0": 2, "x1": 20, "y1": 10}, Decimal("0.91"))], "test-injected")


@pytest.fixture
def processor_storage(tmp_path: Path):
    return LocalStorageService(tmp_path / "storage")


def test_processor_routing_uses_validated_type(processor_storage):
    router = ProcessorRouter(processor_storage)
    assert isinstance(router.get_processor("application/pdf", "pdf"), PDFProcessor)
    assert isinstance(router.get_processor("application/vnd.openxmlformats-officedocument.wordprocessingml.document", "docx"), DOCXProcessor)
    assert isinstance(router.get_processor("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "xlsx"), ExcelProcessor)
    assert isinstance(router.get_processor("image/jpeg", "jpg"), ImageProcessor)
    with pytest.raises(UnsupportedProcessorType):
        router.get_processor("application/pdf", "docx")


def test_pdf_extracts_page_text_and_text_blocks(processor_storage):
    key = "documents/test/pdf"
    save_object(processor_storage, key, make_pdf())
    result = PDFProcessor(processor_storage).process(key, lambda _stage: None)
    assert len(result.pages) == 1
    assert result.pages[0].page_number == 1
    assert "DATA MINE" in result.pages[0].text
    assert result.pages[0].contents
    assert result.pages[0].contents[0].bounding_box is not None


def test_scanned_pdf_detection_calls_ocr(processor_storage):
    document = pymupdf.open()
    document.new_page()
    data = document.tobytes()
    document.close()
    key = "documents/test/scanned-pdf"
    save_object(processor_storage, key, data)
    stages = []
    result = PDFProcessor(processor_storage, ocr=FixedOCR()).process(key, stages.append)
    assert "OCR" in stages
    assert result.pages[0].metadata["ocr_used"] is True
    assert result.pages[0].text == "synthetic OCR line"
    assert result.pages[0].contents[0].confidence == Decimal("0.91")


def test_pdf_table_extraction(processor_storage):
    key = "documents/test/table-pdf"
    save_object(processor_storage, key, make_table_pdf())
    result = PDFProcessor(processor_storage).process(key, lambda _stage: None)
    assert result.pages[0].tables
    assert result.pages[0].tables[0].extraction_method == "pymupdf_find_tables"
    assert any(cell.raw_value == "Alpha" for cell in result.pages[0].tables[0].cells)


def test_ocr_runs_locally_or_reports_tesseract_unavailable(monkeypatch):
    png = BytesIO()
    image = Image.new("RGB", (500, 100), "white")
    ImageDraw.Draw(image).text((10, 25), "DATA MINE", fill="black")
    image.save(png, format="PNG")
    if not shutil.which("tesseract"):
        with pytest.raises(OCREngineUnavailable, match="Tesseract"):
            OCRService().recognize(png.getvalue())
    else:
        result = OCRService().recognize(png.getvalue())
        assert result.method == "tesseract"
        assert result.text


def test_docx_extracts_headings_paragraphs_and_tables(processor_storage):
    document = WordDocument()
    document.add_heading("Geology notes", level=1)
    document.add_paragraph("Synthetic processing fixture")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Column"
    table.cell(0, 1).text = "Value"
    data = BytesIO()
    document.save(data)
    key = "documents/test/file.docx"
    save_object(processor_storage, key, data.getvalue())
    result = DOCXProcessor(processor_storage).process(key, lambda _stage: None)
    assert result.pages[0].page_number is None
    assert result.pages[0].metadata["pagination_available"] is False
    assert any(content.content_type == "heading" for content in result.pages[0].contents)
    assert result.pages[0].tables[0].cells[0].raw_value == "Column"


def test_excel_extracts_sheet_cells_and_preserves_blank_values(processor_storage):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Production"
    sheet["A1"] = "Period"
    sheet["B1"] = "Output"
    sheet["A2"] = "2026-01"
    sheet["B2"] = 12
    sheet["C2"] = None
    payload = BytesIO()
    workbook.save(payload)
    key = "documents/test/workbook.xlsx"
    save_object(processor_storage, key, payload.getvalue())
    result = ExcelProcessor(processor_storage).process(key, lambda _stage: None)
    page = result.pages[0]
    assert page.page_number is None
    assert page.metadata["sheet"] == "Production"
    assert any(cell.raw_value == "12" for cell in page.tables[0].cells)
    assert all(cell.raw_value != "0" for cell in page.tables[0].cells)


def test_legacy_xls_extraction_when_xlrd_is_available(processor_storage):
    xlwt = pytest.importorskip("xlwt")
    workbook = xlwt.Workbook()
    sheet = workbook.add_sheet("Legacy")
    sheet.write(0, 0, "Name")
    sheet.write(0, 1, "Count")
    sheet.write(1, 0, "Sample")
    sheet.write(1, 1, 7)
    payload = BytesIO()
    workbook.save(payload)
    key = "documents/test/legacy.xls"
    save_object(processor_storage, key, payload.getvalue())
    page = ExcelProcessor(processor_storage).process(key, lambda _stage: None).pages[0]
    assert page.metadata["sheet"] == "Legacy"
    assert any(cell.raw_value in {"7", "7.0"} for cell in page.tables[0].cells)


def test_image_processor_records_local_ocr_method(processor_storage):
    image = BytesIO()
    Image.new("RGB", (40, 30), "white").save(image, format="PNG")
    key = "documents/test/image.png"
    save_object(processor_storage, key, image.getvalue())
    result = ImageProcessor(processor_storage, ocr=FixedOCR()).process(key, lambda _stage: None)
    assert result.pages[0].metadata["ocr_method"] == "test-injected"
    assert result.pages[0].contents[0].text == "synthetic OCR line"


@pytest.fixture
def api_client(tmp_path: Path):
    storage = LocalStorageService(tmp_path / "objects")
    async_engine = create_async_engine(DATABASE_URL, poolclass=NullPool)
    async_sessions = async_sessionmaker(async_engine, expire_on_commit=False)

    async def test_db():
        async with async_sessions() as session:
            yield session

    async def test_user():
        return SimpleNamespace(id=uuid4(), username="test-admin", role="ADMIN", is_active=True)

    app.dependency_overrides[get_storage_service] = lambda: storage
    app.dependency_overrides[get_db] = test_db
    app.dependency_overrides[get_current_user] = test_user
    with TestClient(app) as client:
        client.storage = storage
        client.document_ids = []
        yield client

    document_ids = list(client.document_ids)
    sync_engine = __import__("sqlalchemy").create_engine(SYNC_DATABASE_URL)
    try:
        with SyncSession(sync_engine) as session:
            if document_ids:
                ids = [__import__("uuid").UUID(value) for value in document_ids]
                versions = session.execute(select(DocumentVersion).where(DocumentVersion.document_id.in_(ids))).scalars().all()
                docs = session.execute(select(Document).where(Document.id.in_(ids))).scalars().all()
                for key in {item.storage_key for item in [*versions, *docs]}:
                    storage.delete_file(key)
                session.execute(delete(Document).where(Document.id.in_(ids)))
                session.commit()
    finally:
        sync_engine.dispose()
        asyncio.run(async_engine.dispose())
        app.dependency_overrides.clear()


def upload_test_pdf(client: TestClient, suffix: str, content: str = "integration pipeline test text") -> str:
    payload = make_pdf(content + " " + suffix)
    response = client.post("/api/v1/documents/upload", files={"file": (f"processing-{suffix}.pdf", payload, "application/pdf")})
    assert response.status_code == 201, response.text
    document_id = response.json()["id"]
    client.document_ids.append(document_id)
    return document_id


def test_processing_routes_require_authentication(api_client):
    auth_dependency = app.dependency_overrides.pop(get_current_user)
    try:
        document_id = str(uuid4())
        assert api_client.post(f"/api/v1/processing/{document_id}/process").status_code == 401
        assert api_client.get(f"/api/v1/processing/{document_id}/processing-status").status_code == 401
        assert api_client.get(f"/api/v1/processing/{document_id}/pages").status_code == 401
    finally:
        app.dependency_overrides[get_current_user] = auth_dependency


def create_worker_job(client: TestClient, monkeypatch: pytest.MonkeyPatch, document_id: str) -> str:
    import app.api.routes_processing as processing_routes

    monkeypatch.setattr(processing_routes.process_document_task, "apply_async", lambda *args, **kwargs: None)
    response = client.post(f"/api/v1/processing/{document_id}/process")
    assert response.status_code == 202, response.text
    return response.json()["id"]


def test_processing_job_duplicate_prevention_status_and_pages(api_client, monkeypatch):
    document_id = upload_test_pdf(api_client, uuid4().hex)
    job_id = create_worker_job(api_client, monkeypatch, document_id)
    duplicate = api_client.post(f"/api/v1/processing/{document_id}/process")
    assert duplicate.status_code == 409
    status = api_client.get(f"/api/v1/processing/{document_id}/processing-status")
    assert status.status_code == 200
    assert status.json()["job"]["id"] == job_id
    assert status.json()["job"]["status"] == "QUEUED"
    assert api_client.get(f"/api/v1/processing/{document_id}/pages").json() == []


def test_celery_submission_failure_marks_job_failed(api_client, monkeypatch):
    import app.api.routes_processing as processing_routes

    document_id = upload_test_pdf(api_client, uuid4().hex)
    monkeypatch.setattr(processing_routes.process_document_task, "apply_async", lambda *args, **kwargs: (_ for _ in ()).throw(ConnectionError("broker unavailable")))
    response = api_client.post(f"/api/v1/processing/{document_id}/process")
    assert response.status_code == 503
    body = response.json()["detail"]
    assert body["status"] == "FAILED"
    status = api_client.get(f"/api/v1/processing/{document_id}/processing-status").json()
    assert status["job"]["status"] == "FAILED"


def test_worker_success_and_reprocessing_replace_results(api_client, monkeypatch):
    document_id = upload_test_pdf(api_client, uuid4().hex)
    sync_engine = __import__("sqlalchemy").create_engine(SYNC_DATABASE_URL, poolclass=NullPool)
    sessions = sessionmaker(sync_engine, expire_on_commit=False)
    try:
        first_job_id = create_worker_job(api_client, monkeypatch, document_id)
        first_result = process_document_job(first_job_id, api_client.storage, sessions)
        assert first_result["status"] == "COMPLETED"
        first_pages = api_client.get(f"/api/v1/processing/{document_id}/pages")
        assert first_pages.status_code == 200
        assert len(first_pages.json()) == 1
        assert "integration pipeline test text" in first_pages.json()[0]["text"]
        first_counts = None
        with sessions() as session:
            first_counts = (
                session.query(DocumentPage).filter_by(document_version_id=first_pages.json()[0]["document_version_id"]).count(),
                session.query(ExtractedContent).join(DocumentPage).filter(DocumentPage.document_version_id == first_pages.json()[0]["document_version_id"]).count(),
            )
        second_job_id = create_worker_job(api_client, monkeypatch, document_id)
        second_result = process_document_job(second_job_id, api_client.storage, sessions)
        assert second_result["status"] == "COMPLETED"
        with sessions() as session:
            second_counts = (
                session.query(DocumentPage).filter_by(document_version_id=first_pages.json()[0]["document_version_id"]).count(),
                session.query(ExtractedContent).join(DocumentPage).filter(DocumentPage.document_version_id == first_pages.json()[0]["document_version_id"]).count(),
            )
        assert first_counts == second_counts
        status = api_client.get(f"/api/v1/processing/{document_id}/processing-status").json()
        assert status["job"]["status"] == "COMPLETED"
        assert any(stage["stage"] == "VERIFICATION_PENDING" for stage in status["stages"])
    finally:
        sync_engine.dispose()


def test_worker_failure_marks_job_failed(api_client, monkeypatch):
    import app.api.routes_processing as processing_routes

    response = api_client.post("/api/v1/documents/upload", files={"file": (f"broken-{uuid4().hex}.pdf", b"%PDF-1.7\nnot a valid PDF", "application/pdf")})
    assert response.status_code == 201
    document_id = response.json()["id"]
    api_client.document_ids.append(document_id)
    job_id = create_worker_job(api_client, monkeypatch, document_id)
    sync_engine = __import__("sqlalchemy").create_engine(SYNC_DATABASE_URL, poolclass=NullPool)
    try:
        with pytest.raises(Exception):
            process_document_job(job_id, api_client.storage, sessionmaker(sync_engine, expire_on_commit=False))
        status = api_client.get(f"/api/v1/processing/{document_id}/processing-status").json()
        assert status["job"]["status"] == "FAILED"
        assert status["job"]["error_message"]
        assert status["stages"][-1]["status"] == "FAILED"
    finally:
        sync_engine.dispose()
