from datetime import date, datetime
from io import BytesIO
import logging

from app.processors.base import ContentResult, DocumentProcessor, PageResult, ProcessingResult, TableCellResult, TableResult
from app.core.config import MAX_SPREADSHEET_CELLS

logger = logging.getLogger(__name__)


def _as_text(value) -> str | None:
    if value is None:
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


class ExcelProcessor(DocumentProcessor):
    def __init__(self, storage, max_sheet_cells: int = MAX_SPREADSHEET_CELLS):
        super().__init__(storage)
        self.max_sheet_cells = max_sheet_cells

    def process(self, storage_key, stage):
        stage("PARSING")
        data = self._read_document(storage_key)
        if storage_key.lower().endswith(".xls"):
            return self._process_xls(data, stage)
        return self._process_xlsx(data, stage)

    def _process_xlsx(self, data: bytes, stage):
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter

        workbook = load_workbook(BytesIO(data), read_only=True, data_only=False)
        pages = []
        try:
            for sheet_index, worksheet in enumerate(workbook.worksheets):
                stage("PARSING")
                if worksheet.max_row * worksheet.max_column > self.max_sheet_cells:
                    raise ValueError(f"Worksheet {worksheet.title!r} exceeds the supported cell limit")
                contents = []
                cell_results = []
                text_parts = []
                for row_index, row in enumerate(worksheet.iter_rows(), start=0):
                    for column_index, cell in enumerate(row):
                        value = _as_text(cell.value)
                        cell_ref = f"{get_column_letter(column_index + 1)}{row_index + 1}"
                        if value is not None and value != "":
                            contents.append(ContentResult("spreadsheet_cell", value, metadata={"sheet": worksheet.title, "sheet_index": sheet_index, "row_index": row_index, "column_index": column_index, "cell": cell_ref, "data_type": getattr(cell, "data_type", None)}))
                            text_parts.append(value)
                        cell_results.append(TableCellResult(
                            row_index,
                            column_index,
                            value,
                            metadata={"sheet": worksheet.title, "cell": cell_ref, "data_type": getattr(cell, "data_type", None)},
                        ))
                stage("TABLE_EXTRACTION")
                tables = []
                if any(cell.raw_value is not None for cell in cell_results):
                    tables.append(TableResult(1, "openpyxl", metadata={"sheet": worksheet.title, "sheet_index": sheet_index}, cells=cell_results))
                pages.append(PageResult(None, "\n".join(text_parts) or None, {"source_unit": "worksheet", "sheet": worksheet.title, "sheet_index": sheet_index, "pagination_available": False}, contents, tables))
                logger.info("Excel worksheet processed sheet=%s index=%s", worksheet.title, sheet_index)
        finally:
            workbook.close()
        return ProcessingResult(pages)

    def _process_xls(self, data: bytes, stage):
        import xlrd

        workbook = xlrd.open_workbook(file_contents=data, on_demand=True)
        pages = []
        try:
            for sheet_index in range(workbook.nsheets):
                worksheet = workbook.sheet_by_index(sheet_index)
                if worksheet.nrows * worksheet.ncols > self.max_sheet_cells:
                    raise ValueError(f"Worksheet {worksheet.name!r} exceeds the supported cell limit")
                stage("PARSING")
                contents = []
                cells = []
                text_parts = []
                for row_index in range(worksheet.nrows):
                    for column_index in range(worksheet.ncols):
                        cell = worksheet.cell(row_index, column_index)
                        if cell.ctype == xlrd.XL_CELL_EMPTY or cell.ctype == xlrd.XL_CELL_BLANK:
                            value = None
                        elif cell.ctype == xlrd.XL_CELL_DATE:
                            value = xlrd.xldate.xldate_as_datetime(cell.value, workbook.datemode).isoformat()
                        else:
                            value = _as_text(cell.value)
                        metadata = {"sheet": worksheet.name, "sheet_index": sheet_index, "row_index": row_index, "column_index": column_index, "cell_type": cell.ctype}
                        if value is not None and value != "":
                            contents.append(ContentResult("spreadsheet_cell", value, metadata=metadata))
                            text_parts.append(value)
                        cells.append(TableCellResult(row_index, column_index, value, metadata=metadata))
                stage("TABLE_EXTRACTION")
                tables = [TableResult(1, "xlrd", metadata={"sheet": worksheet.name, "sheet_index": sheet_index}, cells=cells)] if any(cell.raw_value is not None for cell in cells) else []
                pages.append(PageResult(None, "\n".join(text_parts) or None, {"source_unit": "worksheet", "sheet": worksheet.name, "sheet_index": sheet_index, "pagination_available": False}, contents, tables))
                logger.info("Excel worksheet processed sheet=%s index=%s", worksheet.name, sheet_index)
        finally:
            workbook.release_resources()
        return ProcessingResult(pages)
