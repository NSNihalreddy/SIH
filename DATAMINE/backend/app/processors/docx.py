from io import BytesIO
import logging

from docx import Document as OpenDocument

from app.processors.base import ContentResult, DocumentProcessor, PageResult, ProcessingResult, TableCellResult, TableResult

logger = logging.getLogger(__name__)


class DOCXProcessor(DocumentProcessor):
    def process(self, storage_key, stage):
        stage("PARSING")
        document = OpenDocument(BytesIO(self._read_document(storage_key)))
        logger.info("DOCX document parsed paragraphs=%s tables=%s", len(document.paragraphs), len(document.tables))
        page_contents = []
        text_parts = []
        for paragraph_index, paragraph in enumerate(document.paragraphs):
            content = paragraph.text
            if not content.strip():
                continue
            style_name = paragraph.style.name if paragraph.style is not None else None
            is_heading = bool(style_name and style_name.lower().startswith("heading"))
            content_type = "heading" if is_heading else "paragraph"
            heading_level = None
            if is_heading:
                suffix = style_name.lower().removeprefix("heading").strip()
                heading_level = int(suffix) if suffix.isdigit() else None
            page_contents.append(ContentResult(content_type, content, metadata={"paragraph_index": paragraph_index, "style": style_name, "heading_level": heading_level}))
            text_parts.append(content)

        stage("TABLE_EXTRACTION")
        tables = []
        for table_index, table in enumerate(document.tables, start=1):
            cells = []
            for row_index, row in enumerate(table.rows):
                for column_index, cell in enumerate(row.cells):
                    value = cell.text
                    cells.append(TableCellResult(
                        row_index,
                        column_index,
                        value if value != "" else None,
                        metadata={"source": "python-docx"},
                    ))
            tables.append(TableResult(table_index, "python-docx", metadata={"row_count": len(table.rows)}, cells=cells))

        page = PageResult(
            page_number=None,
            text="\n".join(text_parts) or None,
            metadata={"source_unit": "document", "pagination_available": False},
            contents=page_contents,
            tables=tables,
        )
        return ProcessingResult([page])
